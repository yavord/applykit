"""Export API tests: base-resume and tailored-revision downloads in PDF/DOCX."""

from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from docx.shared import Pt
from fastapi.testclient import TestClient
from pypdf import PdfReader
from reportlab.pdfgen import canvas

from app.data.models import DocKind
from app.data.repositories import DocumentRepo, JobData, JobRepo
from app.web.main import app
from tests.test_render import _pdf_fontsizes

SAMPLE_LINES = [
    "Jane Q. Developer",
    "Senior Backend Engineer | San Francisco, CA",
    "jane.q@example.com | (415) 555-0132 | linkedin.com/in/janeq",
    "SUMMARY",
    "Experienced backend engineer with 8 years building distributed systems.",
    "SKILLS",
    "Languages: Python, SQL, TypeScript",
    "AWS, Docker, PostgreSQL",
    "EXPERIENCE",
    "Senior Backend Engineer | Acme Corp",
    "2020 - 2023",
    "- Built migration pipeline serving 2M users",
    "- Cut p95 latency by 40%",
    "",
    "Backend Engineer | Globex Inc",
    "2016 - 2020",
    "- Shipped billing rewrite",
    "EDUCATION",
    "State University, 2012 - 2016",
    "B.S. Computer Science",
    "GPA: 3.8",
    "PROJECTS",
    "Resume Parser | personal",
    "2023 - 2024",
    "- Extracts structured resumes from raw text",
    "CERTIFICATIONS",
    "AWS Certified Solutions Architect - Associate, Amazon Web Services, 2024",
]

SENTINEL = "ZQX-Sentinel-Bullet"


def write_pdf(path: Path, lines: list[str]) -> None:
    c = canvas.Canvas(str(path))
    y = 780

    for line in lines:
        c.drawString(60, y, line if line else " ")
        y -= 16

    c.save()


@pytest.fixture
def client():
    return TestClient(app)


@pytest.fixture
def pdf(tmp_path) -> Path:
    path = tmp_path / "resume.pdf"

    write_pdf(path, SAMPLE_LINES)

    return path


def _import(client: TestClient, pdf: Path, name: str) -> dict:
    with pdf.open("rb") as f:
        resp = client.post(
            "/api/resumes/import",
            files={"file": ("resume.pdf", f, "application/pdf")},
            data={"name": name},
        )

    assert resp.status_code == 201, resp.text

    return resp.json()


def _replace_first_bullet(body: dict, text: str) -> dict:
    """Section payload for PATCH with the first experience bullet replaced."""
    for sec in body["sections"]:
        if sec["kind"] == "experience":
            sec["content"][0]["bullets"][0] = text
            break
    else:
        raise AssertionError("fixture resume has no experience section")

    return body


def _patch_sections(client: TestClient, resume_id: int, body: dict) -> dict:
    resp = client.patch(f"/api/resumes/{resume_id}/sections", json={"sections": body["sections"]})
    assert resp.status_code == 200, resp.text

    return resp.json()


def _pdf_text(body: bytes) -> str:
    return "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(body)).pages)


def _docx_text(body: bytes) -> str:
    return "\n".join(p.text for p in Document(BytesIO(body)).paragraphs)


def test_base_export_pdf_reflects_edit(client, pdf):
    imported = _import(client, pdf, "Smith Resume")
    patched = _patch_sections(client, imported["id"], _replace_first_bullet(imported, SENTINEL))

    resp = client.get(f"/api/resumes/{imported['id']}/export.pdf")

    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/pdf"
    assert resp.content.startswith(b"%PDF-")
    assert SENTINEL in _pdf_text(resp.content)
    assert PdfReader(BytesIO(resp.content)).metadata.title == "Smith Resume"
    assert f"-rev{patched['revision']}." in resp.headers["content-disposition"]


def test_base_export_docx_reflects_edit(client, pdf):
    imported = _import(client, pdf, "Smith Resume")
    _patch_sections(client, imported["id"], _replace_first_bullet(imported, SENTINEL))

    resp = client.get(f"/api/resumes/{imported['id']}/export.docx")

    assert resp.status_code == 200
    assert (
        resp.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    text = _docx_text(resp.content)
    assert SENTINEL in text
    assert "Work Experience" in text


def test_marker_unwrapped_in_export(client, pdf):
    """Stored uncertainty markers render as their value; the literal is never leaked."""
    imported = _import(client, pdf, "Smith Resume")

    for sec in imported["sections"]:
        if sec["kind"] == "contact":
            sec["content"]["name"] = {"v": "Urn Ce Tain", "uncertain": True}
            break

    _patch_sections(client, imported["id"], imported)

    resp = client.get(f"/api/resumes/{imported['id']}/export.pdf")
    text = _pdf_text(resp.content)

    assert resp.status_code == 200
    assert "Urn Ce Tain" in text
    assert "uncertain" not in text


def test_revision_export_reflects_revision(client, pdf):
    imported = _import(client, pdf, "Smith Resume")
    job_id = JobRepo().upsert(
        JobData(
            source="test",
            title="Engineer",
            company="Acme",
            norm_company="acme",
            norm_title="engineer",
        )
    )
    doc = DocumentRepo().create(
        imported["id"],
        job_id,
        DocKind.TAILORED_RESUME,
        {
            "sections": [
                {
                    "kind": "experience",
                    "position": 0,
                    "content": [{"title": "Rev", "organization": "Co", "bullets": [SENTINEL]}],
                }
            ]
        },
    )

    for fmt, read_text in (("pdf", _pdf_text), ("docx", _docx_text)):
        resp = client.get(f"/api/documents/{doc.id}/export.{fmt}")

        assert resp.status_code == 200
        assert SENTINEL in read_text(resp.content)
        assert "migration pipeline" not in read_text(resp.content)  # base content absent
        assert f"tailored-{doc.id}." in resp.headers["content-disposition"]


def test_revision_export_missing_404(client):
    resp = client.get("/api/documents/999/export.pdf")

    assert resp.status_code == 404
    assert "no document with id 999" in resp.json()["detail"]


def test_export_unknown_resume_404(client):
    resp = client.get("/api/resumes/999/export.pdf")

    assert resp.status_code == 404
    assert "no resume with id 999" in resp.json()["detail"]


def test_cover_letter_not_exportable_422(client, pdf):
    imported = _import(client, pdf, "Smith Resume")
    job_id = JobRepo().upsert(
        JobData(
            source="test",
            title="Engineer",
            company="Acme",
            norm_company="acme",
            norm_title="engineer",
        )
    )
    doc = DocumentRepo().create(imported["id"], job_id, DocKind.COVER_LETTER, {"text": "hi"})

    resp = client.get(f"/api/documents/{doc.id}/export.pdf")

    assert resp.status_code == 422
    assert "not a tailored resume" in resp.json()["detail"]


def test_malformed_revision_422(client, pdf):
    imported = _import(client, pdf, "Smith Resume")
    job_id = JobRepo().upsert(
        JobData(
            source="test",
            title="Engineer",
            company="Acme",
            norm_company="acme",
            norm_title="engineer",
        )
    )
    doc = DocumentRepo().create(imported["id"], job_id, DocKind.TAILORED_RESUME, {"nope": 1})

    resp = client.get(f"/api/documents/{doc.id}/export.docx")

    assert resp.status_code == 422
    assert "no exportable sections" in resp.json()["detail"]


def test_unknown_format_422(client):
    resp = client.get("/api/resumes/1/export.txt")

    assert resp.status_code == 422


def test_export_params_change_output(client, pdf):
    imported = _import(client, pdf, "Smith Resume")
    rid = imported["id"]

    default = client.get(f"/api/resumes/{rid}/export.pdf")
    with_param = client.get(f"/api/resumes/{rid}/export.pdf?name_size=30")

    assert default.status_code == 200
    assert with_param.status_code == 200
    assert default.content != with_param.content
    assert 24.0 in _pdf_fontsizes(default.content)["Times-Bold"]
    assert 30.0 not in _pdf_fontsizes(default.content)["Times-Bold"]
    assert 30.0 in _pdf_fontsizes(with_param.content)["Times-Bold"]
    assert 24.0 not in _pdf_fontsizes(with_param.content)["Times-Bold"]


def test_export_bad_params_422(client, pdf):
    imported = _import(client, pdf, "Smith Resume")
    rid = imported["id"]

    resp = client.get(f"/api/resumes/{rid}/export.pdf?name_size=99")

    assert resp.status_code == 422
    detail = resp.json()["detail"]
    assert "name_size" in detail
    assert "out of range" in detail

    resp = client.get(f"/api/resumes/{rid}/export.pdf?bogus=1")

    assert resp.status_code == 422
    assert "unknown export setting" in resp.json()["detail"]


def test_export_docx_params(client, pdf):
    imported = _import(client, pdf, "Smith Resume")
    rid = imported["id"]

    default = Document(BytesIO(client.get(f"/api/resumes/{rid}/export.docx").content)).paragraphs[0]
    assert default.runs[0].font.name == "Times New Roman"
    assert default.runs[0].font.size == Pt(24)

    timed = Document(
        BytesIO(
            client.get(
                f"/api/resumes/{rid}/export.docx?font_family=times_new_roman&name_size=28"
            ).content
        )
    ).paragraphs[0]
    assert timed.runs[0].font.name == "Times New Roman"
    assert timed.runs[0].font.size == Pt(28)


def test_revision_export_params(client, pdf):
    imported = _import(client, pdf, "Smith Resume")
    job_id = JobRepo().upsert(
        JobData(
            source="test",
            title="Engineer",
            company="Acme",
            norm_company="acme",
            norm_title="engineer",
        )
    )
    doc = DocumentRepo().create(
        imported["id"],
        job_id,
        DocKind.TAILORED_RESUME,
        {
            "sections": [
                {
                    "kind": "experience",
                    "position": 0,
                    "content": [{"title": "Rev", "organization": "Co", "bullets": [SENTINEL]}],
                }
            ]
        },
    )

    default = client.get(f"/api/documents/{doc.id}/export.docx")
    with_param = client.get(f"/api/documents/{doc.id}/export.docx?font_family=helvetica")

    assert default.status_code == 200
    assert with_param.status_code == 200
    assert default.content != with_param.content
