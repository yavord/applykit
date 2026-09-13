"""Renderer parametrization tests: fonts, sizes, spacing, alignment, margins, layout."""

import re
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Pt
from pypdf import PdfReader

from app.resumes.services.export_service import Block, _layout, render_docx, render_pdf
from app.resumes.services.settings import (
    EducationOrder,
    ExportSettings,
    HeaderAlign,
    SkillsLayout,
)

_SUBSET_TAG = re.compile(r"^[A-Z]{6}\+")


def _font_face(ref) -> str:
    """Face name from a /BaseFont value: drop the leading '/' and any subset tag."""
    return _SUBSET_TAG.sub("", str(ref.get_object()["/BaseFont"]).removeprefix("/"))


def _pdf_fontsizes(body: bytes) -> dict[str, set[float]]:
    """Face -> sizes actually used by drawn text; Tf-only switches draw nothing."""
    page = PdfReader(BytesIO(body)).pages[0]
    faces = {str(tag): _font_face(ref) for tag, ref in page["/Resources"]["/Font"].items()}
    out: dict[str, set[float]] = {}
    current = None
    for operands, op in page.get_contents().operations:
        if op == b"Tf":
            current = (faces[str(operands[0])], round(float(operands[1]), 1))
        elif op in (b"Tj", b"TJ") and current:
            out.setdefault(current[0], set()).add(current[1])
    return out


def _text_starts(body: bytes) -> list[tuple[float, str]]:
    """(x-origin, text) per emitted text piece, page order; tm[4] is the x-origin."""
    starts: list[tuple[float, str]] = []
    for page in PdfReader(BytesIO(body)).pages:
        page.extract_text(
            visitor_text=lambda text, cm, tm, font, size: starts.append((tm[4], text))
        )
    return starts


def _para(document, text: str):
    """First DOCX paragraph whose text matches exactly."""
    return next(p for p in document.paragraphs if p.text == text)


def _blocks() -> list[Block]:
    return [
        Block("h1", "Jane Q. Developer", header=True),
        Block("text", "jane@example.com", header=True),
        Block("h2", "Work Experience"),
        Block("text", "Engineer — Acme Corp", bold=True),
        Block("bullet", "Built the migration pipeline."),
        Block("text", "Plain summary body."),
    ]


def test_pdf_default_fontsizes():
    s = _pdf_fontsizes(render_pdf(_blocks(), "t", ExportSettings()))

    assert s["Times-Bold"] == {24.0, 12.0}
    assert s["Times-Roman"] == {10.0}


def test_pdf_builtin_faces():
    s = _pdf_fontsizes(render_pdf(_blocks(), "t", ExportSettings(font_family="times_new_roman")))

    assert s["Times-Bold"] == {24.0, 12.0}
    assert s["Times-Roman"] == {10.0}

    s = _pdf_fontsizes(render_pdf(_blocks(), "t", ExportSettings(font_family="helvetica")))

    assert s["Helvetica-Bold"] == {24.0, 12.0}
    assert s["Helvetica"] == {10.0}


def test_pdf_name_size_param():
    s = _pdf_fontsizes(
        render_pdf(_blocks(), "t", ExportSettings(font_family="work_sans", name_size=30))
    )

    assert s["WorkSans-Bold"] == {30.0, 12.0}
    assert s["WorkSans-Regular"] == {10.0}


def test_docx_run_fonts_and_sizes():
    d = Document(BytesIO(render_docx(_blocks(), ExportSettings())))

    name = _para(d, "Jane Q. Developer").runs[0]
    assert name.font.name == "Times New Roman"
    assert name.font.size == Pt(24)

    assert _para(d, "Engineer — Acme Corp").runs[0].font.size == Pt(12)
    assert _para(d, "Engineer — Acme Corp").runs[0].bold is True
    assert _para(d, "Plain summary body.").runs[0].font.size == Pt(10)

    d = Document(
        BytesIO(render_docx(_blocks(), ExportSettings(font_family="times_new_roman", name_size=28)))
    )
    name = _para(d, "Jane Q. Developer").runs[0]
    assert name.font.name == "Times New Roman"
    assert name.font.size == Pt(28)


def test_docx_justify():
    d = Document(BytesIO(render_docx(_blocks(), ExportSettings(align_justify=True))))
    assert _para(d, "Plain summary body.").alignment is WD_ALIGN_PARAGRAPH.JUSTIFY
    assert _para(d, "Work Experience").alignment is WD_ALIGN_PARAGRAPH.LEFT

    d = Document(BytesIO(render_docx(_blocks(), ExportSettings(align_justify=False))))
    assert _para(d, "Plain summary body.").alignment is WD_ALIGN_PARAGRAPH.LEFT


def test_pdf_justify_emits_tw():
    blocks = [Block("h1", "t"), Block("text", "word " * 200)]

    def ops_contain_tw(body: bytes) -> bool:
        for page in PdfReader(BytesIO(body)).pages:
            for _, op in page.get_contents().operations:
                if op == b"Tw":
                    return True
        return False

    assert ops_contain_tw(render_pdf(blocks, "t", ExportSettings(align_justify=True)))
    assert not ops_contain_tw(render_pdf(blocks, "t", ExportSettings(align_justify=False)))


def test_docx_header_alignment():
    for align, expected in (
        (HeaderAlign.CENTER, WD_ALIGN_PARAGRAPH.CENTER),
        (HeaderAlign.LEFT, WD_ALIGN_PARAGRAPH.LEFT),
        (HeaderAlign.RIGHT, WD_ALIGN_PARAGRAPH.RIGHT),
    ):
        d = Document(BytesIO(render_docx(_blocks(), ExportSettings(header_align=align))))

        assert _para(d, "Jane Q. Developer").alignment is expected
        assert _para(d, "jane@example.com").alignment is expected


def test_pdf_header_alignment_x():
    xs = {
        a: next(
            x
            for x, t in _text_starts(render_pdf(_blocks(), "t", ExportSettings(header_align=a)))
            if "Jane Q." in t
        )
        for a in HeaderAlign
    }

    assert xs[HeaderAlign.LEFT] < xs[HeaderAlign.CENTER] < xs[HeaderAlign.RIGHT]


def test_pdf_margins_and_spacing_change_page_count():
    blocks = [Block("h1", "Name"), Block("h2", "Section")]
    blocks += [Block("bullet", "word " * 60) for _ in range(30)]

    def pages(settings: ExportSettings) -> int:
        return len(PdfReader(BytesIO(render_pdf(blocks, "t", settings))).pages)

    compact = pages(
        ExportSettings(
            name_size=16,
            header_size=10,
            subheader_size=8,
            body_size=8,
            section_spacing=0,
            entry_spacing=0,
            line_spacing=10,
            margin_top=10,
            margin_bottom=10,
            margin_side=30,
        )
    )
    spacious = pages(
        ExportSettings(
            name_size=30,
            header_size=20,
            subheader_size=18,
            body_size=14,
            section_spacing=10,
            entry_spacing=10,
            line_spacing=15,
            margin_top=50,
            margin_bottom=50,
            margin_side=50,
        )
    )

    assert spacious > compact
    assert spacious >= 2


def test_docx_margins_and_line_spacing():
    d = Document(
        BytesIO(
            render_docx(
                _blocks(),
                ExportSettings(margin_top=50, margin_bottom=50, margin_side=30, line_spacing=15),
            )
        )
    )
    section = d.sections[0]

    assert section.top_margin == Pt(50)
    assert section.bottom_margin == Pt(50)
    assert section.left_margin == Pt(30)
    assert section.right_margin == Pt(30)
    assert _para(d, "Plain summary body.").paragraph_format.line_spacing == 1.5


def _sec(kind, content, position=0):
    return {"kind": kind, "position": position, "content": content}


EDU = _sec("education", [{"degree": "B.S. Computer Science", "institution": "State University"}])
SKILLS = _sec(
    "skills",
    [{"group": "Languages", "skills": ["Python", "SQL"]}, {"group": "", "skills": ["AWS"]}],
)


def test_education_order():
    first = _layout([EDU], ExportSettings(education_order=EducationOrder.DEGREE_FIRST))
    assert first[1] == Block("text", "B.S. Computer Science", bold=True, right="")
    assert first[2] == Block("text", "State University", italic=True)

    second = _layout([EDU], ExportSettings(education_order=EducationOrder.INSTITUTION_FIRST))
    assert second[1] == Block("text", "State University", bold=True, right="")
    assert second[2] == Block("text", "B.S. Computer Science", italic=True)


def test_skills_layouts():
    for layout, expected in (
        (
            SkillsLayout.GROUPED,
            [Block("text", "Python, SQL", label="Languages: "), Block("text", "AWS")],
        ),
        (SkillsLayout.INLINE, [Block("text", "Python, SQL, AWS")]),
        (
            SkillsLayout.COLUMN,
            [Block("text", "Python"), Block("text", "SQL"), Block("text", "AWS")],
        ),
    ):
        blocks = [
            b for b in _layout([SKILLS], ExportSettings(skills_layout=layout)) if b.kind != "h2"
        ]

        assert blocks == expected


def test_pdf_h2_emits_rule():
    body = render_pdf([Block("h1", "t"), Block("h2", "Section")], "t", ExportSettings())

    ops = [
        op for page in PdfReader(BytesIO(body)).pages for _, op in page.get_contents().operations
    ]

    assert b"l" in ops or b"re" in ops


def test_docx_h2_bottom_border():
    d = Document(BytesIO(render_docx(_blocks(), ExportSettings())))

    pBdr = _para(d, "Work Experience")._p.pPr.find(qn("w:pBdr"))

    assert pBdr is not None
    assert pBdr.find(qn("w:bottom")).get(qn("w:val")) == "single"


def test_docx_entry_right_tab_stop():
    d = Document(
        BytesIO(
            render_docx(
                [Block("text", "Acme Corp", bold=True, right="2020 – 2023")], ExportSettings()
            )
        )
    )

    p = d.paragraphs[0]

    assert p.text == "Acme Corp\t2020 – 2023"
    assert p.paragraph_format.tab_stops[0].alignment == WD_TAB_ALIGNMENT.RIGHT
